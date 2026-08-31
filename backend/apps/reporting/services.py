import io
import csv
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from apps.students.models import Student


class DossierService:
    @staticmethod
    def get_full_dossier(student_id):
        student = Student.objects.select_related('user').prefetch_related(
            'semesters',
            'academic_committee__user',
            'tutoring_sessions__created_by',
            'tutoring_sessions__participants__user',
            'tutoring_sessions__observations__autor',
            'agreements__responsable',
            'thesis_progress_records__registrado_por',
            'publications__evidencia',
            'academic_events__evidencia',
            'research_stays__evidencia',
            'other_products__evidencia',
            'evidences__cargado_por'
        ).get(pk=student_id)

        # 1. Datos Demográficos
        demographics = {
            'id': student.id,
            'matricula': student.matricula,
            'nombre_completo': student.nombre_completo,
            'email': student.user.email if student.user else None,
            'programa_doctoral': student.programa_doctoral,
            'fecha_ingreso': str(student.fecha_ingreso),
            'cohorte': student.cohorte,
            'semestre_actual': student.semestre_actual,
            'asesor_principal': student.asesor_principal,
            'coasesor': student.coasesor,
            'estatus_activo': student.estatus_activo
        }

        # 2. Comité Tutoral
        committee = [
            {
                'id': c.id,
                'user_id': c.user_id,
                'nombre': c.user.full_name,
                'email': c.user.email,
                'rol_comite': c.rol_comite,
                'rol_comite_display': c.get_rol_comite_display(),
                'fecha_asignacion': str(c.fecha_asignacion),
                'is_active': c.is_active
            }
            for c in student.academic_committee.all()
        ]

        # 3. Semestres
        semesters = [
            {
                'id': s.id,
                'numero': s.numero,
                'fecha_inicio': str(s.fecha_inicio),
                'fecha_fin': str(s.fecha_fin),
                'is_active': s.is_active
            }
            for s in student.semesters.order_by('numero')
        ]

        # 4. Tutorías
        tutorings = [
            {
                'id': t.id,
                'semester_numero': t.semester.numero,
                'fecha_sesion': str(t.fecha_sesion),
                'modalidad': t.modalidad,
                'resumen_general': t.resumen_general,
                'proxima_reunion_fecha': str(t.proxima_reunion_fecha) if t.proxima_reunion_fecha else None,
                'proxima_reunion_notas': t.proxima_reunion_notas,
                'created_by': t.created_by.full_name if t.created_by else 'Asesor',
                'participantes': [p.user.full_name for p in t.participants.all()],
                'observaciones': [
                    {'tema': o.tema_revisado, 'detalle': o.observaciones_detalladas, 'autor': o.autor.full_name if o.autor else 'Asesor'}
                    for o in t.observations.all()
                ]
            }
            for t in student.tutoring_sessions.order_by('-fecha_sesion')
        ]

        # 5. Acuerdos
        agreements = [
            {
                'id': a.id,
                'semester_numero': a.semester.numero if a.semester else None,
                'descripcion': a.descripcion,
                'responsable_nombre': a.responsable.full_name,
                'fecha_limite': str(a.fecha_limite),
                'estado': a.estado,
                'is_overdue': a.is_overdue
            }
            for a in student.agreements.order_by('-created_at')
        ]

        # 6. Avances de Tesis
        thesis_history = [
            {
                'id': th.id,
                'semester_numero': th.semester.numero,
                'porcentaje_avance': th.porcentaje_avance,
                'fecha_registro': str(th.fecha_registro),
                'componentes_json': th.componentes_json,
                'observaciones': th.observaciones,
                'registrado_por': th.registrado_por.full_name if th.registrado_por else 'Asesor'
            }
            for th in student.thesis_progress_records.order_by('semester__numero')
        ]

        # 7. Publicaciones
        publications = [
            {
                'id': p.id,
                'semester_numero': p.semester.numero,
                'titulo': p.titulo,
                'autores': p.autores_texto,
                'tipo': p.get_tipo_display(),
                'revista_editorial': p.revista_editorial,
                'estado': p.get_estado_display(),
                'fecha_publicacion': str(p.fecha_publicacion) if p.fecha_publicacion else None,
                'doi_url': p.doi_url
            }
            for p in student.publications.order_by('-created_at')
        ]

        # 8. Eventos y Ponencias
        events = [
            {
                'id': ev.id,
                'semester_numero': ev.semester.numero,
                'tipo_evento': ev.get_tipo_evento_display(),
                'nombre_evento': ev.nombre_evento,
                'titulo_ponencia': ev.titulo_ponencia,
                'fecha_presentacion': str(ev.fecha_presentacion),
                'sede_lugar': ev.sede_lugar,
                'modalidad': ev.modalidad
            }
            for ev in student.academic_events.order_by('-fecha_presentacion')
        ]

        # 9. Estancias
        stays = [
            {
                'id': st.id,
                'semester_numero': st.semester.numero,
                'institucion_receptora': st.institucion_receptora,
                'pais': st.pais,
                'fecha_inicio': str(st.fecha_inicio),
                'fecha_fin': str(st.fecha_fin),
                'responsable_estancia': st.responsable_estancia,
                'objetivos': st.objetivos
            }
            for st in student.research_stays.order_by('-fecha_inicio')
        ]

        # 10. Evidencias
        evidences = [
            {
                'id': e.id,
                'semester_numero': e.semester.numero,
                'tipo': e.get_tipo_display(),
                'descripcion': e.descripcion,
                'fecha_carga': str(e.fecha_carga.date()),
                'cargado_por': e.cargado_por.full_name if e.cargado_por else 'Estudiante',
                'archivo_url': e.archivo_adjunto.url if e.archivo_adjunto else None,
                'url_doi': e.url_doi,
                'actividad_tipo': e.actividad_tipo
            }
            for e in student.evidences.order_by('-fecha_carga')
        ]

        return {
            'demographics': demographics,
            'academic_committee': committee,
            'semesters': semesters,
            'tutorings': tutorings,
            'agreements': agreements,
            'thesis_history': thesis_history,
            'publications': publications,
            'academic_events': events,
            'research_stays': stays,
            'evidences': evidences,
            'generated_at': datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        }


class ExportService:
    @staticmethod
    def generate_excel(dossier):
        wb = Workbook()
        ws_info = wb.active
        ws_info.title = "Expediente General"

        # Estilos
        primary_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='6365EF', end_color='6365EF', fill_type='solid')
        title_font = Font(name='Calibri', size=14, bold=True, color='2C1867')
        bold_font = Font(name='Calibri', size=10, bold=True)
        regular_font = Font(name='Calibri', size=10)
        border_side = Side(style='thin', color='D3D3D3')
        thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

        # 1. Hoja: Resumen General
        ws_info.append(["N.E.X.U.S. - REPORTE INTEGRAL DE EXPEDIENTE DOCTORAL"])
        ws_info["A1"].font = title_font
        ws_info.append([])

        demo = dossier['demographics']
        ws_info.append(["Matrícula:", demo['matricula'], "Nombre Completo:", demo['nombre_completo']])
        ws_info.append(["Programa:", demo['programa_doctoral'], "Cohorte:", demo['cohorte']])
        ws_info.append(["Semestre Actual:", f"Semestre {demo['semestre_actual']}", "Fecha de Ingreso:", demo['fecha_ingreso']])
        ws_info.append(["Asesor Principal:", demo['asesor_principal'], "Coasesor:", demo['coasesor'] or 'No asignado'])
        ws_info.append([])

        # Resumen Métricas
        ws_info.append(["MÉTRICAS CONSOLIDADAS DEL EXPEDIENTE"])
        ws_info.merge_cells("A8:D8")
        ws_info["A8"].font = primary_font
        ws_info["A8"].fill = header_fill

        ws_info.append(["Total Sesiones Tutoría:", len(dossier['tutorings']), "Acuerdos Totales:", len(dossier['agreements'])])
        ws_info.append(["Publicaciones Científicas:", len(dossier['publications']), "Ponencias en Congresos:", len(dossier['academic_events'])])
        ws_info.append(["Estancias de Investigación:", len(dossier['research_stays']), "Evidencias Digitales:", len(dossier['evidences'])])
        ws_info.append([])

        # 2. Hoja: Tutorías y Acuerdos
        ws_tut = wb.create_sheet(title="Tutorías y Acuerdos")
        ws_tut.append(["HISTORIAL DE TUTORÍAS"])
        ws_tut["A1"].font = bold_font
        ws_tut.append(["Semestre", "Fecha", "Modalidad", "Registrado Por", "Resumen General", "Próxima Reunión"])
        for col_num in range(1, 7):
            cell = ws_tut.cell(row=2, column=col_num)
            cell.font = primary_font
            cell.fill = header_fill

        for t in dossier['tutorings']:
            ws_tut.append([
                t['semester_numero'],
                t['fecha_sesion'],
                t['modalidad'],
                t['created_by'],
                t['resumen_general'],
                t['proxima_reunion_fecha'] or 'N/A'
            ])

        ws_tut.append([])
        row_agr = ws_tut.max_row + 1
        ws_tut.cell(row=row_agr, column=1, value="LISTADO DE ACUERDOS Y COMPROMISOS").font = bold_font
        ws_tut.append(["ID", "Semestre", "Descripción", "Responsable", "Fecha Límite", "Estado"])
        for col_num in range(1, 7):
            cell = ws_tut.cell(row=row_agr + 1, column=col_num)
            cell.font = primary_font
            cell.fill = header_fill

        for a in dossier['agreements']:
            ws_tut.append([
                f"#{a['id']}",
                a['semester_numero'] or 'N/A',
                a['descripcion'],
                a['responsable_nombre'],
                a['fecha_limite'],
                a['estado']
            ])

        # 3. Hoja: Producción Científica
        ws_prod = wb.create_sheet(title="Producción Científica")
        ws_prod.append(["PUBLICACIONES Y ARTÍCULOS"])
        ws_prod["A1"].font = bold_font
        ws_prod.append(["Sem", "Título", "Autores", "Tipo", "Revista / Editorial", "Estado", "DOI"])
        for col_num in range(1, 8):
            cell = ws_prod.cell(row=2, column=col_num)
            cell.font = primary_font
            cell.fill = header_fill

        for p in dossier['publications']:
            ws_prod.append([
                p['semester_numero'],
                p['titulo'],
                p['autores'],
                p['tipo'],
                p['revista_editorial'],
                p['estado'],
                p['doi_url'] or 'N/A'
            ])

        # Ajuste de ancho de columnas
        for ws in [ws_info, ws_tut, ws_prod]:
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = col[0].column_letter
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    @staticmethod
    def generate_pdf(dossier):
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        primary_color = colors.HexColor('#6365EF')
        dark_color = colors.HexColor('#2C1867')
        light_bg = colors.HexColor('#F8F9FC')

        title_style = ParagraphStyle(
            'NexusTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=16,
            leading=18,
            textColor=dark_color
        )
        subtitle_style = ParagraphStyle(
            'NexusSubtitle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            textColor=colors.HexColor('#667085')
        )
        section_heading = ParagraphStyle(
            'NexusSectionHeading',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=11,
            leading=14,
            textColor=primary_color,
            spaceAfter=6
        )
        table_cell_bold = ParagraphStyle(
            'CellBold',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=8,
            leading=10,
            textColor=colors.black
        )
        table_cell = ParagraphStyle(
            'CellRegular',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8,
            leading=10,
            textColor=colors.black
        )

        story = []
        demo = dossier['demographics']

        # Header Institucional
        story.append(Paragraph("N.E.X.U.S. — NÚCLEO DE EXPEDIENTE Y SEGUIMIENTO UNIVERSITARIO", title_style))
        story.append(Paragraph(f"Cédula Oficial de Expediente Doctoral • Generado el {dossier['generated_at']}", subtitle_style))
        story.append(Spacer(1, 10))
        story.append(HRFlowable(width="100%", thickness=1.5, color=primary_color, spaceBefore=2, spaceAfter=10))

        # 1. Datos Generales del Doctorando
        story.append(Paragraph("1. DATOS GENERALES Y TRAYECTORIA ACADÉMICA", section_heading))
        demo_table_data = [
            [
                Paragraph("<b>Doctorando:</b>", table_cell),
                Paragraph(demo['nombre_completo'], table_cell_bold),
                Paragraph("<b>Matrícula:</b>", table_cell),
                Paragraph(demo['matricula'], table_cell_bold)
            ],
            [
                Paragraph("<b>Programa:</b>", table_cell),
                Paragraph(demo['programa_doctoral'], table_cell),
                Paragraph("<b>Cohorte:</b>", table_cell),
                Paragraph(demo['cohorte'], table_cell)
            ],
            [
                Paragraph("<b>Semestre Actual:</b>", table_cell),
                Paragraph(f"Semestre {demo['semestre_actual']}", table_cell),
                Paragraph("<b>Fecha Ingreso:</b>", table_cell),
                Paragraph(demo['fecha_ingreso'], table_cell)
            ],
            [
                Paragraph("<b>Asesor Principal:</b>", table_cell),
                Paragraph(demo['asesor_principal'], table_cell),
                Paragraph("<b>Coasesor:</b>", table_cell),
                Paragraph(demo['coasesor'] or 'No asignado', table_cell)
            ]
        ]
        demo_table = Table(demo_table_data, colWidths=[90, 180, 80, 190])
        demo_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), light_bg),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E4E7EC')),
            ('PADDING', (0, 0), (-1, -1), 5),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        story.append(demo_table)
        story.append(Spacer(1, 12))

        # 2. Historial de Tutorías Recientes
        story.append(Paragraph("2. SESIONES DE TUTORÍA Y SEGUIMIENTO ACADÉMICO", section_heading))
        tut_table_data = [[
            Paragraph("<b>Sem</b>", table_cell_bold),
            Paragraph("<b>Fecha</b>", table_cell_bold),
            Paragraph("<b>Modalidad</b>", table_cell_bold),
            Paragraph("<b>Asesor / Evaluador</b>", table_cell_bold),
            Paragraph("<b>Resumen General</b>", table_cell_bold)
        ]]
        for t in dossier['tutorings'][:6]:
            tut_table_data.append([
                Paragraph(str(t['semester_numero']), table_cell),
                Paragraph(t['fecha_sesion'], table_cell),
                Paragraph(t['modalidad'], table_cell),
                Paragraph(t['created_by'], table_cell),
                Paragraph(t['resumen_general'][:100] + ('...' if len(t['resumen_general']) > 100 else ''), table_cell)
            ])
        if len(dossier['tutorings']) == 0:
            tut_table_data.append([Paragraph("Sin sesiones registradas", table_cell), "", "", "", ""])

        tut_table = Table(tut_table_data, colWidths=[30, 60, 65, 110, 275])
        tut_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#EEF0FF')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E4E7EC')),
            ('PADDING', (0, 0), (-1, -1), 4),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ]))
        story.append(tut_table)
        story.append(Spacer(1, 12))

        # 3. Acuerdos y Compromisos
        story.append(Paragraph("3. ACUERDOS Y ESTADO DE COMPROMISOS", section_heading))
        agr_table_data = [[
            Paragraph("<b>ID</b>", table_cell_bold),
            Paragraph("<b>Compromiso</b>", table_cell_bold),
            Paragraph("<b>Responsable</b>", table_cell_bold),
            Paragraph("<b>Fecha Límite</b>", table_cell_bold),
            Paragraph("<b>Estado</b>", table_cell_bold)
        ]]
        for a in dossier['agreements'][:8]:
            agr_table_data.append([
                Paragraph(f"#{a['id']}", table_cell),
                Paragraph(a['descripcion'][:90] + ('...' if len(a['descripcion']) > 90 else ''), table_cell),
                Paragraph(a['responsable_nombre'], table_cell),
                Paragraph(a['fecha_limite'], table_cell),
                Paragraph(a['estado'], table_cell_bold)
            ])
        if len(dossier['agreements']) == 0:
            agr_table_data.append([Paragraph("Sin acuerdos registrados", table_cell), "", "", "", ""])

        agr_table = Table(agr_table_data, colWidths=[30, 250, 110, 70, 80])
        agr_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#EEF0FF')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E4E7EC')),
            ('PADDING', (0, 0), (-1, -1), 4),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        story.append(agr_table)
        story.append(Spacer(1, 12))

        # 4. Producción Científica
        story.append(Paragraph("4. PRODUCCIÓN CIENTÍFICA, CONGRESOS Y ESTANCIAS", section_heading))
        pub_table_data = [[
            Paragraph("<b>Tipo</b>", table_cell_bold),
            Paragraph("<b>Título de la Contribución</b>", table_cell_bold),
            Paragraph("<b>Revista / Congreso / Sede</b>", table_cell_bold),
            Paragraph("<b>Estado</b>", table_cell_bold)
        ]]
        for p in dossier['publications']:
            pub_table_data.append([
                Paragraph(p['tipo'], table_cell),
                Paragraph(p['titulo'][:80], table_cell),
                Paragraph(p['revista_editorial'], table_cell),
                Paragraph(p['estado'], table_cell)
            ])
        for ev in dossier['academic_events']:
            pub_table_data.append([
                Paragraph(ev['tipo_evento'], table_cell),
                Paragraph(ev['titulo_ponencia'][:80], table_cell),
                Paragraph(f"{ev['nombre_evento']} ({ev['sede_lugar']})", table_cell),
                Paragraph("PRESENTADO", table_cell)
            ])
        if len(dossier['publications']) == 0 and len(dossier['academic_events']) == 0:
            pub_table_data.append([Paragraph("Sin productos científicos registrados", table_cell), "", "", ""])

        pub_table = Table(pub_table_data, colWidths=[90, 220, 150, 80])
        pub_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#EEF0FF')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E4E7EC')),
            ('PADDING', (0, 0), (-1, -1), 4),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ]))
        story.append(pub_table)

        doc.build(story)
        buffer.seek(0)
        return buffer.getvalue()
